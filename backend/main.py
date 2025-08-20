from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from dotenv import load_dotenv
import google.generativeai as genai
import os
import uvicorn
import traceback

from docx import Document   # ✅ for DOCX export
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from .story_manager import (
    create_new_section,
    get_section_content,
    append_to_section,
    get_all_sections,
    delete_section,
    get_next_chapter_number,
    get_last_chapter_number,
)


# ==========================================================
# ✅ Setup
# ==========================================================
load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    raise RuntimeError("❌ GEMINI_API_KEY is missing. Add it to your .env file.")

genai.configure(api_key=GEMINI_API_KEY)

# 🔥 Switch model here
model = genai.GenerativeModel("gemini-2.0-flash")   # instead of "gemini-1.5-flash-latest"

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory="frontend"), name="static")


# ==========================================================
# ✅ Endpoints
# ==========================================================
@app.get("/")
async def serve_index():
    return FileResponse("frontend/index.html")


@app.get("/sections")
async def sections():
    return {"sections": get_all_sections()}


@app.get("/get_section_content/{section}")
async def get_section(section: str):
    content = get_section_content(section)
    if content is None:
        raise HTTPException(status_code=404, detail="Section not found")
    return {"content": content}


@app.post("/start_new_section/")
async def start_new_section(data: dict):
    name = data.get("name")
    if not name:
        raise HTTPException(status_code=400, detail="Section name required")
    file = create_new_section(name)
    return {"message": f"Section '{name}' created.", "filepath": file}


@app.delete("/delete_section/{section}")
async def remove_section(section: str):
    ok = delete_section(section)
    if not ok:
        raise HTTPException(status_code=404, detail="Section not found")
    return {"message": f"Section '{section}' deleted."}


@app.post("/generate_next_part/")
async def next_part(data: dict):
    section = data.get("section")
    user_prompt = data.get("user_prompt")

    if not section or not user_prompt:
        raise HTTPException(status_code=400, detail="Section and prompt required")

    content = get_section_content(section)
    if content is None:
        raise HTTPException(status_code=404, detail="Section not found")

    # ✅ Decide whether to revise or continue
    if "revise" in user_prompt.lower():
        chap_num = get_last_chapter_number(content)
        heading = f"### Chapter {chap_num} (Revised):"
    else:
        chap_num = get_next_chapter_number(content)
        heading = f"### Chapter {chap_num}:"

    full_prompt = (
        f"Here is the story so far:\n\n{content}\n\n"
        f"Task: {user_prompt}\n\n"
        f"Continue with proper heading '{heading} ...'. "
        "Always keep consistent style and narrative flow."
    )

    try:
        response = model.generate_content(full_prompt)

        # ✅ Extract text safely
        new_text = getattr(response, "text", None)
        if not new_text and hasattr(response, "candidates"):
            parts = response.candidates[0].content.parts
            new_text = "".join(p.text for p in parts if hasattr(p, "text"))

        if not new_text:
            raise Exception("No text returned from Gemini model")

        new_text = new_text.strip()
        append_to_section(section, f"{heading}\n\n{new_text}")
        return {"new_story_part": f"{heading}\n\n{new_text}"}

    except Exception as e:
        traceback.print_exc()  # ✅ Show full error in console
        raise HTTPException(status_code=500, detail=f"Gemini error: {str(e)}")


# ==========================================================
# ✅ Download DOCX & PDF
# ==========================================================
@app.get("/download_docx/{section}")
async def download_docx(section: str):
    content = get_section_content(section)
    if content is None:
        raise HTTPException(status_code=404, detail="Section not found")

    safe_name = section.replace(" ", "_")
    filepath = f"{safe_name}.docx"

    doc = Document()
    doc.add_heading(section, level=1)
    for line in content.splitlines():
        doc.add_paragraph(line)
    doc.save(filepath)

    return FileResponse(
        filepath,
        filename=filepath,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


@app.get("/download_pdf/{section}")
async def download_pdf(section: str):
    content = get_section_content(section)
    if content is None:
        raise HTTPException(status_code=404, detail="Section not found")

    safe_name = section.replace(" ", "_")
    filepath = f"{safe_name}.pdf"

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(filepath)
    story = [Paragraph(section, styles['Title'])]

    for line in content.splitlines():
        story.append(Paragraph(line, styles['Normal']))

    doc.build(story)

    return FileResponse(filepath, filename=filepath, media_type="application/pdf")


# ==========================================================
# ✅ Run
# ==========================================================
if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
